import pytest
from app.core.config import get_settings


@pytest.fixture
def signed_in(client, monkeypatch, tmp_path):
    monkeypatch.setattr(get_settings(), 'upload_root', str(tmp_path))
    assert client.post('/api/auth/login', json={'email': 'admin@example.com', 'password': 'Admin123!'}).status_code == 200
    return client


def test_folder_crud_and_nonempty_protection(signed_in):
    client = signed_in
    folder = client.post('/api/documents/folders', json={'name': 'Finance'}).json()
    fid = folder['folder_id']
    assert client.get('/api/documents/folders').json()[0]['name'] == 'Finance'
    assert client.patch(f'/api/documents/folders/{fid}', json={'name': 'Reports'}).json()['name'] == 'Reports'
    upload = client.post(f'/api/documents/folders/{fid}/upload', files={'file': ('report.csv', b'name,revenue\nNorth,120', 'text/csv')})
    assert upload.status_code == 200
    did = upload.json()['document_id']
    document = client.get('/api/documents').json()[0]
    assert document['folder_id'] == fid
    assert document['size'] == len(b'name,revenue\nNorth,120')
    assert client.get('/api/documents/folders').json()[0]['document_count'] == 1
    assert client.delete(f'/api/documents/folders/{fid}').status_code == 409
    assert client.patch(f'/api/documents/{did}', json={'folder_id': None}).status_code == 200
    assert client.delete(f'/api/documents/folders/{fid}').status_code == 200
    assert client.get('/api/documents/folders').json() == []
    assert client.delete(f'/api/documents/{did}').status_code == 200
    assert client.get('/api/documents').json() == []


def test_folder_validation_and_access(signed_in):
    client = signed_in
    assert client.post('/api/documents/folders', json={'name': '   '}).status_code == 422
    folder = client.post('/api/documents/folders', json={'name': 'Private'}).json()
    assert client.post('/api/documents/folders', json={'name': 'Private'}).status_code == 409
    client.post('/api/auth/login', json={'email': 'manager@example.com', 'password': 'Manager123!'})
    fid = folder['folder_id']
    assert client.get('/api/documents/folders').json() == []
    assert client.delete(f'/api/documents/folders/{fid}').status_code == 404
    assert client.post(f'/api/documents/folders/{fid}/upload', files={'file': ('x.txt', b'Hello', 'text/plain')}).status_code == 404


def test_upload_formats_and_errors(signed_in):
    assert signed_in.post('/api/documents/upload', files={'file': ('empty.txt', b'', 'text/plain')}).status_code == 400
    assert signed_in.post('/api/documents/upload', files={'file': ('bad.exe', b'unsupported', 'application/octet-stream')}).status_code == 400
    assert signed_in.post('/api/documents/upload', files={'file': ('bad.pdf', b'not a pdf', 'application/pdf')}).status_code in (400, 422)
    response = signed_in.post('/api/documents/upload', files={'file': ('notes.txt', b'Business notes', 'text/plain')})
    assert response.status_code == 200
    assert signed_in.get('/api/documents').json()[0]['size'] == 14

@pytest.mark.parametrize('kind', ['docx', 'xlsx', 'pdf'])
def test_existing_binary_formats(signed_in, kind):
    from io import BytesIO
    content = BytesIO()
    if kind == 'docx':
        from docx import Document
        document = Document()
        document.add_paragraph('Business document')
        document.save(content)
    elif kind == 'xlsx':
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(['Revenue', 120])
        workbook.save(content)
    else:
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        writer.write(content)
    response = signed_in.post('/api/documents/upload', files={'file': (f'business.{kind}', content.getvalue(), 'application/octet-stream')})
    assert response.status_code == 200
    assert signed_in.get('/api/documents').json()[0]['size'] == len(content.getvalue())
