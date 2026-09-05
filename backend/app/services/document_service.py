from __future__ import annotations
from pathlib import Path
from uuid import uuid4
import csv, io
from pypdf import PdfReader
from docx import Document as DocxDocument
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from app.core.config import get_settings
from app.db.models import Document,DocumentChunk

class DocumentService:
    def extract_text(self,filename:str,content:bytes)->str:
        ext=Path(filename).suffix.lower()
        if ext==".pdf": return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(content)).pages)
        if ext==".docx": return "\n".join(p.text for p in DocxDocument(io.BytesIO(content)).paragraphs)
        if ext in {".csv",".txt",".md",".html"}: return content.decode("utf-8",errors="replace")
        if ext in {".xlsx",".xlsm"}:
            wb=load_workbook(io.BytesIO(content),read_only=True,data_only=True); lines=[]
            for ws in wb.worksheets:
                lines.append(f"# Sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True): lines.append(" | ".join("" if v is None else str(v) for v in row))
            return "\n".join(lines)
        raise ValueError("Unsupported file type")
    def chunk(self,text:str,size:int=900,overlap:int=120)->list[str]:
        text=" ".join(text.split()); out=[]; start=0
        while start<len(text):
            out.append(text[start:start+size]); start += max(1,size-overlap)
        return out or [""]
    def ingest(self,db:Session,filename:str,content:bytes,mime_type:str="application/octet-stream",allowed_roles:str|None=None)->str:
        text=self.extract_text(filename,content); did=f"doc_{uuid4().hex}"
        path=Path(get_settings().upload_root)/f"{did}_{Path(filename).name}"; path.write_bytes(content)
        db.add(Document(id=did,filename=filename,mime_type=mime_type,allowed_roles=allowed_roles or "admin,executive,manager,analyst,employee"))
        for i,ch in enumerate(self.chunk(text)): db.add(DocumentChunk(document_id=did,chunk_index=i,text=ch))
        db.commit(); return did
    def compare(self,left:str,right:str)->dict:
        a=set(x.strip() for x in left.splitlines() if x.strip()); b=set(x.strip() for x in right.splitlines() if x.strip())
        return {"only_left":sorted(a-b),"only_right":sorted(b-a),"shared_count":len(a&b)}
document_service=DocumentService()
